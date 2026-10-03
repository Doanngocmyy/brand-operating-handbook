"""Build templates/SOP-Tracker.xlsx: order tracker (SOP-A), case tracker (SOP-B), SLA flags, KPI dashboard.

Usage: python tools/build_tracker.py   (needs openpyxl)
No customer personal data columns by design: the order number is the only key.
"""
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "templates" / "SOP-Tracker.xlsx"
N_DON, N_CASE = 400, 200
F = "Arial"
INK, MUTED, LINE = "1D1D1F", "6E6E73", "E5E5E7"
ROLE = {"CX": "2F6FB0", "OPS": "6B5A45", "SRC": "2E7D4F", "LOG": "0F7C80", "FIN": "7A4FA3", "CEO": "1D1D1F", "SYS": "8E8E93"}
INPUT = PatternFill("solid", fgColor="FFF9E6")
HEAD = PatternFill("solid", fgColor="F3EFE9")
RED = PatternFill("solid", fgColor="FDECEA")
AMBER = PatternFill("solid", fgColor="FFF4DC")
GREEN = PatternFill("solid", fgColor="E6F4EA")
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(**k):
    return Font(name=F, **k)


wb = Workbook()

# ------------------------------------------------------------------ DS (lists)
ds = wb.active
ds.title = "DS"
STATUSES = [("10 Mới", "CX", "Đã thanh toán"), ("20 Đã xác nhận", "SRC", "Email xác nhận + ETA"),
            ("30 Đã đặt NCC", "SRC", "Có mã PO"), ("40 NCC đã gửi", "SRC", "Vận đơn nội địa"),
            ("50 QC đạt", "LOG", "Ảnh QC kho gom"), ("60 Đã xuất", "LOG", "Lượt quét thật"),
            ("70 Thông quan", "LOG", "Tới cảng / hải quan"), ("80 Đang giao", "LOG", "Hãng giao nhận hàng"),
            ("90 Đã giao", "CX", "POD / ảnh giao"), ("95 Hủy/hoàn", "–", "Đã hoàn tiền"), ("99 Đóng", "–", "D+14 không case")]
CASES = [("C1 Hủy đơn", "B1"), ("C2 Trễ hạn", "B1"), ("C3 Hư hỏng", "B1·B2"), ("C4 Thiếu kiện/phụ kiện", "B1"),
         ("C5 Sai hàng/khác mô tả", "B1·B2"), ("C6 Đổi ý", "B2"), ("C7 Chargeback", "B3"), ("C8 Bảo hành", "B1")]
SOLS = ["R1 Sửa", "R2 Gửi bù", "R3 Đổi mới", "R4 Hoàn một phần", "R4 Hoàn 100%", "R4 Tín dụng", "R5 Thu hồi", "Không bồi thường"]
RCS = ["RC-NCC", "RC-PACK", "RC-QC", "RC-FWD", "RC-LAST", "RC-LIST", "RC-CX", "RC-CUST"]
APPR = ["CX", "OPS", "CEO"]
MARKETS = [("AU", 8, 12), ("SG", 5, 8), ("NZ", 10, 14)]
CHANNELS = ["Email", "Form web", "Messenger", "Instagram", "WhatsApp", "Điện thoại", "Cổng thanh toán"]
cols = {"A": ("Trạng thái", [s[0] for s in STATUSES]), "B": ("Chủ trì", [s[1] for s in STATUSES]),
        "C": ("Bằng chứng để vào", [s[2] for s in STATUSES]), "E": ("Loại case", [c[0] for c in CASES]),
        "F": ("Luồng", [c[1] for c in CASES]), "G": ("Giải pháp", SOLS), "H": ("Cấp duyệt", APPR),
        "I": ("Nguyên nhân gốc", RCS), "K": ("Thị trường", [m[0] for m in MARKETS]),
        "L": ("ETA min (tuần)", [m[1] for m in MARKETS]), "M": ("ETA max (tuần)", [m[2] for m in MARKETS]),
        "O": ("Kênh", CHANNELS)}
for c, (h, vals) in cols.items():
    ds[f"{c}1"] = h
    ds[f"{c}1"].font = font(bold=True)
    ds[f"{c}1"].fill = HEAD
    for i, v in enumerate(vals, start=2):
        ds[f"{c}{i}"] = v
        ds[f"{c}{i}"].font = font()
    ds.column_dimensions[c].width = 24
ds["K6"] = "Khung ETA = chính sách đã công bố (chỉnh tại L2:M4). NZ chưa bán, để sẵn."
ds["K6"].font = font(italic=True, color=MUTED)
ds["H6"] = "Mức duyệt: ≤ A$50 CX · 51–150 OPS · > 150 hoặc hoàn 100% CEO"
ds["H6"].font = font(italic=True, color=MUTED)

# ------------------------------------------------------------------ DON (orders)
don = wb.create_sheet("DON", 0)
DON_COLS = [  # header, width, kind (in=input, f=formula), owner
    ("Mã đơn", 12, "in", "CX"), ("Thị trường", 9, "in", "CX"), ("SKU", 16, "in", "CX"), ("Giá trị (A$)", 11, "in", "CX"),
    ("Ngày thanh toán", 12, "in", "CX"), ("ETA min", 11, "f", "SYS"), ("ETA max (đã hứa)", 12, "f", "SYS"),
    ("20 Xác nhận", 11, "in", "CX"), ("30 Đặt NCC", 11, "in", "SRC"), ("Hẹn NCC gửi", 11, "in", "SRC"),
    ("40 NCC gửi", 11, "in", "SRC"), ("50 QC đạt", 11, "in", "SRC"), ("60 Xuất (quét thật)", 12, "in", "LOG"),
    ("70 Thông quan", 11, "in", "LOG"), ("80 Bàn giao giao cuối", 12, "in", "LOG"), ("90 Đã giao (POD)", 12, "in", "LOG"),
    ("99 Đóng", 11, "in", "OPS"), ("95 Hủy/hoàn", 11, "in", "FIN"), ("Cập nhật KH gần nhất", 12, "in", "CX"),
    ("Trạng thái", 16, "f", "SYS"), ("Chủ trì hiện tại", 10, "f", "SYS"), ("Ngày ở trạng thái", 10, "f", "SYS"),
    ("⚑ Chưa đặt NCC >24h", 11, "f", "SYS"), ("⚑ NCC quá hẹn >2d", 11, "f", "SYS"), ("⚑ Sắp vượt ETA (≤14d)", 11, "f", "SYS"),
    ("⚑ Vượt ETA", 9, "f", "SYS"), ("⚑ >7d chưa cập nhật KH", 11, "f", "SYS"), ("Giao trong khung", 9, "f", "SYS"),
    ("Lead time tổng (ngày)", 10, "f", "SYS"), ("Số case", 8, "f", "SYS"),
    ("S1–S2 TT→Đặt NCC", 10, "f", "SYS"), ("S3 Đặt→NCC gửi", 10, "f", "SYS"), ("S4 Gửi→QC đạt", 10, "f", "SYS"),
    ("S5 QC→Xuất", 10, "f", "SYS"), ("S6–S7 Xuất→Giao", 10, "f", "SYS"),
]
for i, (h, w, kind, owner) in enumerate(DON_COLS, start=1):
    c = don.cell(row=2, column=i, value=h)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ROLE[owner])
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = BOX
    don.column_dimensions[L(i)].width = w
don.row_dimensions[2].height = 44
don["A1"] = "SOP-A · Theo dõi đơn  —  ô vàng: nhập ngày khi có BẰNG CHỨNG · cột xám: tự tính · màu tiêu đề = vị trí nhập"
don["A1"].font = font(bold=True, size=12)
don.freeze_panes = "B3"

for r in range(3, N_DON + 3):
    f = {
        "F": f'=IF($E{r}="","",$E{r}+7*IFERROR(VLOOKUP($B{r},DS!$K$2:$M$4,2,FALSE),12))',
        "G": f'=IF($E{r}="","",$E{r}+7*IFERROR(VLOOKUP($B{r},DS!$K$2:$M$4,3,FALSE),12))',
        "T": (f'=IF($A{r}="","",IF($R{r}<>"","95 Hủy/hoàn",IF($Q{r}<>"","99 Đóng",IF($P{r}<>"","90 Đã giao",'
              f'IF($O{r}<>"","80 Đang giao",IF($N{r}<>"","70 Thông quan",IF($M{r}<>"","60 Đã xuất",IF($L{r}<>"","50 QC đạt",'
              f'IF($K{r}<>"","40 NCC đã gửi",IF($I{r}<>"","30 Đã đặt NCC",IF($H{r}<>"","20 Đã xác nhận","10 Mới")))))))))))'),
        "U": f'=IF($T{r}="","",IFERROR(VLOOKUP($T{r},DS!$A$2:$B$12,2,FALSE),""))',
        "V": f'=IF(OR($T{r}="",LEFT($T{r},2)="99",LEFT($T{r},2)="95"),"",TODAY()-MAX($E{r},$H{r}:$P{r}))',
        "W": f'=IF(AND($E{r}<>"",$I{r}="",$R{r}=""),IF(TODAY()-$E{r}>1,"⚑",""),"")',
        "X": f'=IF(AND($J{r}<>"",$K{r}="",$R{r}=""),IF(TODAY()-$J{r}>2,"⚑",""),"")',
        "Y": f'=IF(AND($G{r}<>"",$M{r}="",$R{r}=""),IF(AND($G{r}-TODAY()<=14,$G{r}>=TODAY()),"⚑",""),"")',
        "Z": f'=IF(AND($G{r}<>"",$P{r}="",$R{r}=""),IF(TODAY()>$G{r},"⚑",""),"")',
        "AA": f'=IF(AND($E{r}<>"",$P{r}="",$R{r}=""),IF(TODAY()-MAX($E{r},$H{r},$M{r},$S{r})>7,"⚑",""),"")',
        "AB": f'=IF(OR($P{r}="",$G{r}=""),"",IF($P{r}<=$G{r},"Có","Không"))',
        "AC": f'=IF(OR($P{r}="",$E{r}=""),"",$P{r}-$E{r})',
        "AD": f'=IF($A{r}="","",COUNTIF(CASE!$B$3:$B${N_CASE + 2},$A{r}))',
        "AE": f'=IF(OR($I{r}="",$E{r}=""),"",$I{r}-$E{r})',
        "AF": f'=IF(OR($K{r}="",$I{r}=""),"",$K{r}-$I{r})',
        "AG": f'=IF(OR($L{r}="",$K{r}=""),"",$L{r}-$K{r})',
        "AH": f'=IF(OR($M{r}="",$L{r}=""),"",$M{r}-$L{r})',
        "AI": f'=IF(OR($P{r}="",$M{r}=""),"",$P{r}-$M{r})',
    }
    for i, (h, w, kind, owner) in enumerate(DON_COLS, start=1):
        col = L(i)
        c = don[f"{col}{r}"]
        c.font = font(size=10)
        c.border = BOX
        if col in f:
            c.value = f[col]
            c.fill = PatternFill("solid", fgColor="F5F5F4")
        else:
            c.fill = INPUT
        if col in ("E", "F", "G") or (8 <= i <= 19):
            c.number_format = "dd/mm/yy"
        if col == "D":
            c.number_format = "#,##0"
        if col in ("W", "X", "Y", "Z", "AA", "AB", "AD", "V") or i >= 29:
            c.alignment = Alignment(horizontal="center")

last = N_DON + 2
dv = DataValidation(type="list", formula1="=DS!$K$2:$K$4", allow_blank=True)
don.add_data_validation(dv)
dv.add(f"B3:B{last}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2025,1,1)", allow_blank=True,
                     error="Nhập ngày (dd/mm/yy)", errorTitle="Sai định dạng")
don.add_data_validation(dvd)
dvd.add(f"E3:E{last}")
dvd.add(f"H3:S{last}")
for col in ("W", "X", "Y", "Z", "AA"):
    don.conditional_formatting.add(f"{col}3:{col}{last}", CellIsRule(operator="equal", formula=['"⚑"'], fill=RED, font=Font(color="A5281B", bold=True)))
don.conditional_formatting.add(f"AB3:AB{last}", CellIsRule(operator="equal", formula=['"Có"'], fill=GREEN))
don.conditional_formatting.add(f"AB3:AB{last}", CellIsRule(operator="equal", formula=['"Không"'], fill=RED))
don.conditional_formatting.add(f"T3:T{last}", FormulaRule(formula=[f'LEFT($T3,2)="95"'], fill=RED))
don.conditional_formatting.add(f"T3:T{last}", FormulaRule(formula=[f'LEFT($T3,2)="99"'], fill=GREEN))
don.conditional_formatting.add(f"AD3:AD{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=AMBER))
don.auto_filter.ref = f"A2:{L(len(DON_COLS))}{last}"

# ------------------------------------------------------------------ CASE
cs = wb.create_sheet("CASE", 1)
CASE_COLS = [
    ("Mã case", 11, "in", "CX"), ("Mã đơn", 11, "in", "CX"), ("Loại", 20, "in", "CX"), ("Kênh", 12, "in", "CX"),
    ("Mở lúc", 15, "in", "CX"), ("Phản hồi đầu lúc", 15, "in", "CX"), ("Đủ bằng chứng", 11, "in", "CX"),
    ("Giải pháp", 16, "in", "CX"), ("Đền bù / hoàn (A$)", 11, "in", "CX"), ("Cấp duyệt cần", 9, "f", "SYS"),
    ("Người duyệt", 9, "in", "OPS"), ("Đúng cấp?", 8, "f", "SYS"), ("Thực thi xong", 11, "in", "SRC"),
    ("Ngày hoàn tiền", 11, "in", "FIN"), ("Đóng case", 11, "in", "CX"), ("Nguyên nhân gốc", 11, "in", "OPS"),
    ("Thu hồi từ", 12, "in", "SRC"), ("Số thu hồi (A$)", 10, "in", "FIN"), ("Chi phí ròng (A$)", 10, "f", "SYS"),
    ("Giờ phản hồi", 9, "f", "SYS"), ("Số ngày mở", 9, "f", "SYS"), ("⚑ Phản hồi >24h", 10, "f", "SYS"),
    ("⚑ Mở >7 ngày", 10, "f", "SYS"), ("⚑ Duyệt sai cấp", 10, "f", "SYS"), ("Ghi chú (không ghi dữ liệu cá nhân)", 36, "in", "CX"),
]
for i, (h, w, kind, owner) in enumerate(CASE_COLS, start=1):
    c = cs.cell(row=2, column=i, value=h)
    c.font = font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=ROLE[owner])
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    c.border = BOX
    cs.column_dimensions[L(i)].width = w
cs.row_dimensions[2].height = 44
cs["A1"] = "SOP-B · Theo dõi case hậu mãi  —  mỗi case một dòng · đóng case phải có nguyên nhân gốc"
cs["A1"].font = font(bold=True, size=12)
cs.freeze_panes = "C3"
lastc = N_CASE + 2
for r in range(3, lastc + 1):
    f = {
        "J": f'=IF($I{r}="","",IF(OR($I{r}>150,$H{r}="R4 Hoàn 100%"),"CEO",IF($I{r}>50,"OPS","CX")))',
        "L": f'=IF(OR($J{r}="",$K{r}=""),"",IF(MATCH($K{r},DS!$H$2:$H$4,0)>=MATCH($J{r},DS!$H$2:$H$4,0),"✓","✗"))',
        "S": f'=IF($A{r}="","",N($I{r})-N($R{r}))',
        "T": f'=IF(OR($E{r}="",$F{r}=""),"",ROUND(($F{r}-$E{r})*24,1))',
        "U": f'=IF($E{r}="","",INT(IF($O{r}="",NOW(),$O{r})-$E{r}))',
        "V": f'=IF($E{r}="","",IF($F{r}="",IF((NOW()-$E{r})*24>24,"⚑",""),IF(($F{r}-$E{r})*24>24,"⚑","")))',
        "W": f'=IF(OR($E{r}="",$O{r}<>""),"",IF(NOW()-$E{r}>7,"⚑",""))',
        "X": f'=IF($L{r}="✗","⚑","")',
    }
    for i, (h, w, kind, owner) in enumerate(CASE_COLS, start=1):
        col = L(i)
        c = cs[f"{col}{r}"]
        c.font = font(size=10)
        c.border = BOX
        if col in f:
            c.value = f[col]
            c.fill = PatternFill("solid", fgColor="F5F5F4")
            c.alignment = Alignment(horizontal="center")
        else:
            c.fill = INPUT
        if col in ("E", "F"):
            c.number_format = "dd/mm/yy hh:mm"
        if col in ("G", "M", "N", "O"):
            c.number_format = "dd/mm/yy"
        if col in ("I", "R", "S"):
            c.number_format = "#,##0"
for col, rng in (("C", "=DS!$E$2:$E$9"), ("D", "=DS!$O$2:$O$8"), ("H", "=DS!$G$2:$G$9"), ("K", "=DS!$H$2:$H$4"), ("P", "=DS!$I$2:$I$9")):
    v = DataValidation(type="list", formula1=rng, allow_blank=True)
    cs.add_data_validation(v)
    v.add(f"{col}3:{col}{lastc}")
for col in ("V", "W", "X"):
    cs.conditional_formatting.add(f"{col}3:{col}{lastc}", CellIsRule(operator="equal", formula=['"⚑"'], fill=RED, font=Font(color="A5281B", bold=True)))
cs.conditional_formatting.add(f"L3:L{lastc}", CellIsRule(operator="equal", formula=['"✗"'], fill=RED))
cs.conditional_formatting.add(f"L3:L{lastc}", CellIsRule(operator="equal", formula=['"✓"'], fill=GREEN))
cs.auto_filter.ref = f"A2:{L(len(CASE_COLS))}{lastc}"

# ------------------------------------------------------------------ DASHBOARD
db = wb.create_sheet("DASHBOARD", 0)
db.sheet_view.showGridLines = False
db["B2"] = "Bảng điều khiển vận hành"
db["B2"].font = font(bold=True, size=18)
db["B3"] = '="Cập nhật theo ngày: "&TEXT(TODAY(),"dd/mm/yyyy")&"  ·  dữ liệu từ tab DON và CASE"'
db["B3"].font = font(color=MUTED)
D, C = f"DON!$A$3:$A${last}", f"CASE!$A$3:$A${lastc}"
tiles = [
    ("Đơn đang chạy", f'=COUNTIFS(DON!$T$3:$T${last},"<>",DON!$T$3:$T${last},"<>99 Đóng",DON!$T$3:$T${last},"<>95 Hủy/hoàn")-COUNTIF(DON!$T$3:$T${last},"")+COUNTBLANK(DON!$T$3:$T${last})-COUNTBLANK(DON!$T$3:$T${last})', "0", None),
    ("Giao trong khung", f'=IFERROR(COUNTIF(DON!$AB$3:$AB${last},"Có")/(COUNTIF(DON!$AB$3:$AB${last},"Có")+COUNTIF(DON!$AB$3:$AB${last},"Không")),"–")', "0%", "≥ 90%"),
    ("Đặt NCC ≤ 24h", f'=IFERROR(COUNTIFS(DON!$AE$3:$AE${last},"<=1")/COUNT(DON!$AE$3:$AE${last}),"–")', "0%", "≥ 95%"),
    ("Cờ đỏ đơn đang mở", f'=COUNTIF(DON!$W$3:$AA${last},"⚑")', "0", "0"),
    ("Case / đơn", f'=IFERROR(COUNTA({C})/COUNTA({D}),"–")', "0.0%", "≤ 3%"),
    ("Phản hồi đầu (giờ, trung vị)", f'=IFERROR(MEDIAN(CASE!$T$3:$T${lastc}),"–")', "0.0", "≤ 4"),
    ("Chargeback / đơn", f'=IFERROR(COUNTIF(CASE!$C$3:$C${lastc},"C7*")/COUNTA({D}),"–")', "0.0%", "≤ 0,5%"),
    ("Chi phí case ròng (A$)", f'=SUM(CASE!$S$3:$S${lastc})', "#,##0", None),
]
# simpler, clean formula for running orders
tiles[0] = ("Đơn đang chạy", f'=COUNTA({D})-COUNTIF(DON!$T$3:$T${last},"99*")-COUNTIF(DON!$T$3:$T${last},"95*")', "0", None)
for k, (label, formula, fmt, target) in enumerate(tiles):
    col = 2 + (k % 4) * 3
    row = 5 + (k // 4) * 5
    for rr in range(row, row + 4):
        for cc in range(col, col + 2):
            db.cell(row=rr, column=cc).fill = PatternFill("solid", fgColor="FBFBFA")
            db.cell(row=rr, column=cc).border = Border(top=thin if rr == row else None, bottom=thin if rr == row + 3 else None,
                                                       left=thin if cc == col else None, right=thin if cc == col + 1 else None)
    db.cell(row=row, column=col, value=label).font = font(color=MUTED, size=10)
    v = db.cell(row=row + 1, column=col, value=formula)
    v.font = font(bold=True, size=22)
    v.number_format = fmt
    v.alignment = Alignment(horizontal="left")
    if target:
        db.cell(row=row + 3, column=col, value=f"Mục tiêu {target}").font = font(color=MUTED, size=9)
for c in range(2, 14):
    db.column_dimensions[L(c)].width = 12
db.column_dimensions["A"].width = 3


def table(top, left, title, rows, formula_fn, header=("Nhóm", "Số")):
    db.cell(row=top, column=left, value=title).font = font(bold=True, size=12)
    db.cell(row=top + 1, column=left, value=header[0]).font = font(bold=True, color=MUTED, size=10)
    db.cell(row=top + 1, column=left + 1, value=header[1]).font = font(bold=True, color=MUTED, size=10)
    for i, lab in enumerate(rows):
        r = top + 2 + i
        db.cell(row=r, column=left, value=lab).font = font(size=10)
        c = db.cell(row=r, column=left + 1, value=formula_fn(lab))
        c.font = font(size=10)
        c.number_format = "0.0" if header[1] != "Số" else "0"
    return top + 2, top + 1 + len(rows)


def chart(title, cats_rng, vals_rng, anchor, color, ytitle=None):
    ch = BarChart()
    ch.type = "bar"
    ch.style = 10
    ch.title = title
    ch.legend = None
    ch.height, ch.width = 7.2, 13
    ch.add_data(vals_rng, titles_from_data=False)
    ch.set_categories(cats_rng)
    ch.series[0].graphicalProperties.solidFill = color
    ch.series[0].graphicalProperties.line.solidFill = color
    ch.y_axis.majorGridlines = None
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    if ytitle:
        ch.y_axis.title = ytitle
    db.add_chart(ch, anchor)


# data tables live in columns P:Q (right side); charts on the left
T = 16  # column P
r1, r2 = table(5, T, "Đơn theo trạng thái", [s[0] for s in STATUSES],
               lambda lab: f'=COUNTIF(DON!$T$3:$T${last},"{lab}")')
chart("Đơn theo trạng thái (đang ở đâu)", Reference(db, min_col=T, min_row=r1, max_row=r2),
      Reference(db, min_col=T + 1, min_row=r1, max_row=r2), "B16", "6B5A45")
lt = ["S1–S2 TT→Đặt NCC", "S3 Đặt→NCC gửi", "S4 Gửi→QC đạt", "S5 QC→Xuất", "S6–S7 Xuất→Giao"]
ltcol = {"S1–S2 TT→Đặt NCC": "AE", "S3 Đặt→NCC gửi": "AF", "S4 Gửi→QC đạt": "AG", "S5 QC→Xuất": "AH", "S6–S7 Xuất→Giao": "AI"}
r3, r4 = table(19, T, "Lead time TB theo giai đoạn (ngày)", lt,
               lambda lab: f'=IFERROR(AVERAGE(DON!${ltcol[lab]}$3:${ltcol[lab]}${last}),0)', ("Giai đoạn", "Ngày TB"))
chart("Lead time trung bình theo giai đoạn (ngày)", Reference(db, min_col=T, min_row=r3, max_row=r4),
      Reference(db, min_col=T + 1, min_row=r3, max_row=r4), "H16", "0F7C80")
flags = [("Chưa đặt NCC >24h", "W"), ("NCC quá hẹn >2d", "X"), ("Sắp vượt ETA", "Y"), ("Vượt ETA", "Z"), (">7d chưa cập nhật KH", "AA"),
         ("Case phản hồi >24h", "CV"), ("Case mở >7 ngày", "CW"), ("Case duyệt sai cấp", "CX")]
fl = {a: b for a, b in flags}
r5, r6 = table(27, T, "Cờ SLA đang bật", [a for a, _ in flags],
               lambda lab: (f'=COUNTIF(DON!${fl[lab]}$3:${fl[lab]}${last},"⚑")' if not fl[lab].startswith("C")
                            else f'=COUNTIF(CASE!${fl[lab][1:]}$3:${fl[lab][1:]}${lastc},"⚑")'))
chart("Cờ SLA đang bật → xử lý trong họp 09:00", Reference(db, min_col=T, min_row=r5, max_row=r6),
      Reference(db, min_col=T + 1, min_row=r5, max_row=r6), "B31", "C0392B")
r7, r8 = table(38, T, "Case theo loại", [c[0] for c in CASES],
               lambda lab: f'=COUNTIF(CASE!$C$3:$C${lastc},"{lab}")')
chart("Case theo loại", Reference(db, min_col=T, min_row=r7, max_row=r8),
      Reference(db, min_col=T + 1, min_row=r7, max_row=r8), "H31", "2F6FB0")
r9, r10 = table(49, T, "Case theo nguyên nhân gốc", RCS, lambda lab: f'=COUNTIF(CASE!$P$3:$P${lastc},"{lab}")')
chart("Nguyên nhân gốc → ai sửa quy trình", Reference(db, min_col=T, min_row=r9, max_row=r10),
      Reference(db, min_col=T + 1, min_row=r9, max_row=r10), "B46", "2E7D4F")
db.column_dimensions[L(T)].width = 24
db.page_setup.orientation = "landscape"
db.sheet_properties.pageSetUpPr.fitToPage = True
db.page_setup.fitToWidth, db.page_setup.fitToHeight = 1, 1
db.column_dimensions[L(T + 1)].width = 9

# ------------------------------------------------------------------ HUONG DAN
hd = wb.create_sheet("HUONG_DAN", 0)
hd.sheet_view.showGridLines = False
hd.column_dimensions["A"].width = 3
hd.column_dimensions["B"].width = 26
hd.column_dimensions["C"].width = 90
lines = [
    ("SOP Tracker · [BRAND]", None, "title"),
    ("Đi kèm sổ tay: SOP-A Đơn chuẩn, SOP-B Hậu mãi, KPI & tracking.", None, "muted"),
    ("", None, None),
    ("Cách dùng", None, "h"),
    ("1. Ô vàng", "Nhập tay. Chỉ nhập ngày khi đã có BẰNG CHỨNG (mã PO, ảnh QC, lượt quét thật, POD).", None),
    ("2. Ô xám", "Tự tính: trạng thái, chủ trì hiện tại, cờ SLA, lead time. Không gõ đè.", None),
    ("3. Màu tiêu đề", "Màu = vị trí nhập cột đó: xanh dương CX · xanh lá SRC · xanh ngọc LOG · nâu OPS · tím FIN · xám tự tính.", None),
    ("4. Họp 09:00", "Mở DASHBOARD → biểu đồ 'Cờ SLA đang bật'. Lọc từng cột ⚑ ở DON/CASE, giao người xử lý ngay.", None),
    ("5. Thứ Sáu", "OPS chụp DASHBOARD vào báo cáo tuần. Đếm case theo nguyên nhân gốc, quyết định tắt SKU/NCC.", None),
    ("6. Dữ liệu", "KHÔNG ghi tên, email, SĐT, địa chỉ khách. Mã đơn là khóa nối. Thông tin khách chỉ nằm trong Shopify.", None),
    ("7. Phân quyền", "Khi có nhiều người: tách bản cho từng vị trí hoặc dùng bản xem lọc. Tab chi phí và NCC để file riêng.", None),
    ("8. Dòng DEMO", "Dòng có mã DEMO-… chỉ để xem cách hoạt động. Xóa trước khi dùng thật.", None),
    ("", None, None),
    ("Cờ SLA", None, "h"),
    ("⚑ Chưa đặt NCC >24h", "Đã thanh toán > 1 ngày mà chưa có ngày Đặt NCC → SRC", None),
    ("⚑ NCC quá hẹn >2d", "Quá ngày hẹn NCC gửi 2 ngày chưa có vận đơn → SRC nhắc; +5d báo OPS; +7d đổi NCC", None),
    ("⚑ Sắp vượt ETA", "Còn ≤ 14 ngày tới ETA max mà chưa xuất → CX báo khách trước ≥ 7 ngày (case C2)", None),
    ("⚑ Vượt ETA", "Quá ETA max chưa giao → CX chủ động đề nghị tín dụng 5% hoặc hủy hoàn 100%", None),
    ("⚑ >7d chưa cập nhật KH", "Lần liên hệ gần nhất > 7 ngày → CX gửi cập nhật tuần, ghi ngày vào cột 'Cập nhật KH gần nhất'", None),
    ("⚑ Case phản hồi >24h", "Chưa có người thật trả lời sau 24h → CX, lên đầu họp 09:00", None),
    ("⚑ Case mở >7 ngày", "Case chưa đóng sau 7 ngày → OPS gọi khách", None),
    ("⚑ Duyệt sai cấp", "Người duyệt thấp hơn mức cần (≤50 CX · 51–150 OPS · >150 / hoàn 100% CEO) → OPS rà", None),
]
for i, (a, b, kind) in enumerate(lines, start=2):
    ca = hd.cell(row=i, column=2, value=a)
    if kind == "title":
        ca.font = font(bold=True, size=18)
    elif kind == "muted":
        ca.font = font(color=MUTED)
    elif kind == "h":
        ca.font = font(bold=True, size=13)
    else:
        ca.font = font(bold=True, size=10)
    if b:
        cb = hd.cell(row=i, column=3, value=b)
        cb.font = font(size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")

# ------------------------------------------------------------------ DEMO rows
T0 = date(2026, 10, 4)
def d(n):
    return T0 + timedelta(days=n) if n is not None else None

demo = [  # id, mkt, sku, value, paid, conf, po, promise, sup, qc, ship, cust, hand, deliv, close, cancel, lastupd
    ("DEMO-1001", "AU", "TVC-OAK-180", 980, -95, -95, -94, -80, -79, -74, -70, -45, -36, -33, -19, None, -40),
    ("DEMO-1002", "AU", "CFT-WAL-120", 720, -88, -88, -86, -73, -71, -67, -63, -40, -34, -30, None, None, -35),
    ("DEMO-1003", "SG", "SHO-WHT-100", 540, -45, -45, -44, -32, -31, -27, -21, -6, -3, -1, None, None, -6),
    ("DEMO-1004", "AU", "BUF-BLK-160", 1250, -70, -70, -69, -55, -50, -46, -41, -12, None, None, None, None, -12),
    ("DEMO-1005", "AU", "BKS-OAK-080", 610, -60, -60, -59, -45, -44, -40, -35, None, None, None, None, None, -10),
    ("DEMO-1006", "AU", "DTB-WAL-180", 1490, -40, -40, -39, -24, -15, -12, None, None, None, None, None, None, -9),
    ("DEMO-1007", "SG", "TVC-WHT-160", 830, -20, -20, -19, -6, None, None, None, None, None, None, None, None, -6),
    ("DEMO-1008", "AU", "CFT-OAK-100", 690, -12, -12, -11, -3, None, None, None, None, None, None, None, None, -12),
    ("DEMO-1009", "AU", "SHO-GRY-120", 560, -3, -3, None, None, None, None, None, None, None, None, None, None, -3),
    ("DEMO-1010", "AU", "BUF-OAK-140", 1100, -1, -1, None, None, None, None, None, None, None, None, None, None, -1),
    ("DEMO-1011", "AU", "DTB-OAK-160", 1320, -80, -80, -79, -65, -64, -60, -56, None, None, None, None, None, -8),
    ("DEMO-1012", "AU", "BKS-BLK-100", 650, -30, -30, None, None, None, None, None, None, None, None, None, -29, -29),
]
for i, row in enumerate(demo):
    r = 3 + i
    (oid, mkt, sku, val, paid, conf, po, prom, sup, qc, ship, cust, hand, deliv, close, cancel, lu) = row
    don[f"A{r}"], don[f"B{r}"], don[f"C{r}"], don[f"D{r}"] = oid, mkt, sku, val
    for col, n in (("E", paid), ("H", conf), ("I", po), ("J", prom), ("K", sup), ("L", qc), ("M", ship), ("N", cust),
                   ("O", hand), ("P", deliv), ("Q", close), ("R", cancel), ("S", lu)):
        if n is not None:
            don[f"{col}{r}"] = d(n)
casedemo = [  # id, order, type, channel, open(days, hour), first(days, hour), evid, sol, amt, appr, done, refund, close, rc, from, rec
    ("DEMO-C01", "DEMO-1003", "C4 Thiếu kiện/phụ kiện", "Email", (-1, 9), (-1, 11), -1, "R2 Gửi bù", 15, "CX", None, None, None, "RC-PACK", "NCC", None),
    ("DEMO-C02", "DEMO-1002", "C3 Hư hỏng", "Form web", (-29, 8), (-29, 10), -28, "R4 Hoàn một phần", 72, "OPS", -27, -27, -26, "RC-LAST", "Hãng giao", 72),
    ("DEMO-C03", "DEMO-1012", "C1 Hủy đơn", "Email", (-30, 14), (-30, 15), -30, "R4 Hoàn 100%", 650, "CEO", -29, -29, -29, "RC-CUST", None, None),
    ("DEMO-C04", "DEMO-1011", "C2 Trễ hạn", "Messenger", (-9, 10), (-8, 16), -9, "R4 Tín dụng", 49, "CX", None, None, None, "RC-NCC", None, None),
    ("DEMO-C05", "DEMO-1001", "C8 Bảo hành", "WhatsApp", (-10, 9), (-10, 9.5), -9, "R1 Sửa", 120, "CX", -6, None, -5, "RC-NCC", "NCC", 120),
]
from datetime import datetime
for i, row in enumerate(casedemo):
    r = 3 + i
    (cid, oid, typ, ch, op, fr, ev, sol, amt, appr, done, ref, close, rc, frm, rec) = row
    cs[f"A{r}"], cs[f"B{r}"], cs[f"C{r}"], cs[f"D{r}"] = cid, oid, typ, ch
    cs[f"E{r}"] = datetime.combine(d(op[0]), datetime.min.time()) + timedelta(hours=op[1])
    cs[f"F{r}"] = datetime.combine(d(fr[0]), datetime.min.time()) + timedelta(hours=fr[1])
    cs[f"G{r}"] = d(ev)
    cs[f"H{r}"], cs[f"I{r}"], cs[f"K{r}"] = sol, amt, appr
    for col, n in (("M", done), ("N", ref), ("O", close)):
        if n is not None:
            cs[f"{col}{r}"] = d(n)
    cs[f"P{r}"] = rc
    if frm:
        cs[f"Q{r}"] = frm
    if rec:
        cs[f"R{r}"] = rec

wb.move_sheet("DS", offset=0)
wb._sheets = [wb["HUONG_DAN"], wb["DASHBOARD"], wb["DON"], wb["CASE"], wb["DS"]]
wb.active = 1
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print("saved", OUT)
