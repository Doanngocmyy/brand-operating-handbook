"""Build the LIVE Google Sheets version of the tracker: "[BRAND] Source of Truth".

Usage: python tools/build_sot_sheet.py out.xlsx   → upload to Google Drive with conversion to Google Sheets.
Differences from templates/SOP-Tracker.xlsx (the Excel template):
- Google-native: one ARRAYFORMULA per calculated column (row 3), so new rows calculate automatically.
- No demo rows. Extra tabs: SOURCE_OF_TRUTH, MASTER_SKU, LAUNCH_BLOCKERS (data from site_map.py).
- No customer personal data columns; supplier identity and FIN costs stay in separate restricted files.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import site_map as SM  # noqa: E402

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else "Source_of_Truth.xlsx")
F = "Arial"
ROLE = {"CX": "2F6FB0", "OPS": "6B5A45", "SRC": "2E7D4F", "LOG": "0F7C80", "FIN": "7A4FA3", "SYS": "8E8E93", "INK": "1D1D1F"}
RED = PatternFill("solid", fgColor="FDECEA")
GREEN = PatternFill("solid", fgColor="E6F4EA")
AMBER = PatternFill("solid", fgColor="FFF4DC")
MAXR = 1000  # validation / formatting extent


def font(**k):
    return Font(name=F, **k)


def header(ws, cols, title, row=2):
    ws["A1"] = title
    ws["A1"].font = font(bold=True, size=12)
    for i, (h, w, owner) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=ROLE[owner])
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = "B3"


wb = Workbook()

# ------------------------------------------------------------------ DS
ds = wb.active
ds.title = "DS"
STATUSES = [("10 Mới", "CX"), ("20 Đã xác nhận", "SRC"), ("30 Đã đặt NCC", "SRC"), ("40 NCC đã gửi", "SRC"),
            ("50 QC đạt", "LOG"), ("60 Đã xuất", "LOG"), ("70 Thông quan", "LOG"), ("80 Đang giao", "LOG"),
            ("90 Đã giao", "CX"), ("95 Hủy/hoàn", "–"), ("99 Đóng", "–")]
CASES = ["C1 Hủy đơn", "C2 Trễ hạn", "C3 Hư hỏng", "C4 Thiếu kiện/phụ kiện", "C5 Sai hàng/khác mô tả", "C6 Đổi ý", "C7 Chargeback", "C8 Bảo hành"]
SOLS = ["R1 Sửa", "R2 Gửi bù", "R3 Đổi mới", "R4 Hoàn một phần", "R4 Hoàn 100%", "R4 Tín dụng", "R5 Thu hồi", "Không bồi thường"]
RCS = ["RC-NCC", "RC-PACK", "RC-QC", "RC-FWD", "RC-LAST", "RC-LIST", "RC-CX", "RC-CUST"]
APPR = ["CX", "OPS", "CEO"]
MARKETS = [("AU", 8, 12), ("SG", 5, 8), ("NZ", 10, 14)]
CHANNELS = ["Email", "Form web", "Messenger", "Instagram", "WhatsApp", "Điện thoại", "Cổng thanh toán"]
lists = {"A": ("Trạng thái", [s for s, _ in STATUSES]), "B": ("Chủ trì", [o for _, o in STATUSES]),
         "E": ("Loại case", CASES), "G": ("Giải pháp", SOLS), "H": ("Cấp duyệt", APPR), "I": ("Nguyên nhân gốc", RCS),
         "K": ("Thị trường", [m for m, _, _ in MARKETS]), "L": ("ETA min (tuần)", [a for _, a, _ in MARKETS]),
         "M": ("ETA max (tuần)", [b for _, _, b in MARKETS]), "O": ("Kênh", CHANNELS)}
for c, (h, vals) in lists.items():
    ds[f"{c}1"] = h
    ds[f"{c}1"].font = font(bold=True)
    for i, v in enumerate(vals, start=2):
        ds[f"{c}{i}"] = v
    ds.column_dimensions[c].width = 22
ds["K6"] = "Khung ETA = chính sách đã công bố (sửa ở L2:M4)."

# ------------------------------------------------------------------ DON
don = wb.create_sheet("DON")
DON = [("Mã đơn", 12, "CX"), ("Thị trường", 9, "CX"), ("SKU", 16, "CX"), ("Giá trị (A$)", 10, "CX"), ("Ngày thanh toán", 11, "CX"),
       ("ETA min", 10, "SYS"), ("ETA max (đã hứa)", 11, "SYS"), ("20 Xác nhận", 10, "CX"), ("30 Đặt NCC", 10, "SRC"),
       ("Hẹn NCC gửi", 10, "SRC"), ("40 NCC gửi", 10, "SRC"), ("50 QC đạt", 10, "SRC"), ("60 Xuất (quét thật)", 11, "LOG"),
       ("70 Thông quan", 10, "LOG"), ("80 Bàn giao giao cuối", 11, "LOG"), ("90 Đã giao (POD)", 11, "LOG"), ("99 Đóng", 10, "OPS"),
       ("95 Hủy/hoàn", 10, "FIN"), ("Cập nhật KH gần nhất", 11, "CX"), ("Trạng thái", 15, "SYS"), ("Chủ trì hiện tại", 9, "SYS"),
       ("Ngày ở trạng thái", 9, "SYS"), ("⚑ Chưa đặt NCC >24h", 10, "SYS"), ("⚑ NCC quá hẹn >2d", 10, "SYS"),
       ("⚑ Sắp vượt ETA (≤14d)", 10, "SYS"), ("⚑ Vượt ETA", 9, "SYS"), ("⚑ >7d chưa cập nhật KH", 10, "SYS"),
       ("Giao trong khung", 9, "SYS"), ("Lead time tổng (ngày)", 9, "SYS"), ("Số case", 7, "SYS"),
       ("S1–S2 TT→Đặt NCC", 9, "SYS"), ("S3 Đặt→NCC gửi", 9, "SYS"), ("S4 Gửi→QC đạt", 9, "SYS"), ("S5 QC→Xuất", 9, "SYS"),
       ("S6–S7 Xuất→Giao", 9, "SYS")]
header(don, DON, "SOP-A · Theo dõi đơn — nhập ngày khi có BẰNG CHỨNG · cột xám tự tính (không gõ đè) · màu tiêu đề = vị trí nhập")
status_date = ('IF(R3:R<>"",R3:R,IF(Q3:Q<>"",Q3:Q,IF(P3:P<>"",P3:P,IF(O3:O<>"",O3:O,IF(N3:N<>"",N3:N,IF(M3:M<>"",M3:M,'
               'IF(L3:L<>"",L3:L,IF(K3:K<>"",K3:K,IF(I3:I<>"",I3:I,IF(H3:H<>"",H3:H,E3:E))))))))))')
DONF = {
    "F": '=ARRAYFORMULA(IF(E3:E="","",E3:E+7*IFERROR(VLOOKUP(B3:B,DS!$K$2:$M$4,2,FALSE),12)))',
    "G": '=ARRAYFORMULA(IF(E3:E="","",E3:E+7*IFERROR(VLOOKUP(B3:B,DS!$K$2:$M$4,3,FALSE),12)))',
    "T": ('=ARRAYFORMULA(IF(A3:A="","",IF(R3:R<>"","95 Hủy/hoàn",IF(Q3:Q<>"","99 Đóng",IF(P3:P<>"","90 Đã giao",IF(O3:O<>"","80 Đang giao",'
          'IF(N3:N<>"","70 Thông quan",IF(M3:M<>"","60 Đã xuất",IF(L3:L<>"","50 QC đạt",IF(K3:K<>"","40 NCC đã gửi",'
          'IF(I3:I<>"","30 Đã đặt NCC",IF(H3:H<>"","20 Đã xác nhận","10 Mới"))))))))))))'),
    "U": '=ARRAYFORMULA(IF(T3:T="","",IFERROR(VLOOKUP(T3:T,DS!$A$2:$B$12,2,FALSE),"")))',
    "V": f'=ARRAYFORMULA(IF((T3:T="")+(LEFT(T3:T,2)="99")+(LEFT(T3:T,2)="95"),"",TODAY()-{status_date}))',
    "W": '=ARRAYFORMULA(IF((E3:E<>"")*(I3:I="")*(R3:R=""),IF(TODAY()-E3:E>1,"⚑",""),""))',
    "X": '=ARRAYFORMULA(IF((J3:J<>"")*(K3:K="")*(R3:R=""),IF(TODAY()-J3:J>2,"⚑",""),""))',
    "Y": '=ARRAYFORMULA(IF((G3:G<>"")*(M3:M="")*(R3:R=""),IF((G3:G-TODAY()<=14)*(G3:G>=TODAY()),"⚑",""),""))',
    "Z": '=ARRAYFORMULA(IF((G3:G<>"")*(P3:P="")*(R3:R=""),IF(TODAY()>G3:G,"⚑",""),""))',
    "AA": ('=ARRAYFORMULA(IF((E3:E<>"")*(P3:P="")*(R3:R=""),IF(TODAY()-BYROW({E3:E,H3:H,M3:M,S3:S},LAMBDA(r,MAX(r)))>7,"⚑",""),""))'),
    "AB": '=ARRAYFORMULA(IF((P3:P="")+(G3:G=""),"",IF(P3:P<=G3:G,"Có","Không")))',
    "AC": '=ARRAYFORMULA(IF((P3:P="")+(E3:E=""),"",P3:P-E3:E))',
    "AD": '=ARRAYFORMULA(IF(A3:A="","",COUNTIF(CASE!B3:B,A3:A)))',
    "AE": '=ARRAYFORMULA(IF((I3:I="")+(E3:E=""),"",I3:I-E3:E))',
    "AF": '=ARRAYFORMULA(IF((K3:K="")+(I3:I=""),"",K3:K-I3:I))',
    "AG": '=ARRAYFORMULA(IF((L3:L="")+(K3:K=""),"",L3:L-K3:K))',
    "AH": '=ARRAYFORMULA(IF((M3:M="")+(L3:L=""),"",M3:M-L3:L))',
    "AI": '=ARRAYFORMULA(IF((P3:P="")+(M3:M=""),"",P3:P-M3:M))',
}
for col, f in DONF.items():
    don[f"{col}3"] = f
dv = DataValidation(type="list", formula1="=DS!$K$2:$K$4", allow_blank=True)
don.add_data_validation(dv)
dv.add(f"B3:B{MAXR}")
for col in ("W", "X", "Y", "Z", "AA"):
    don.conditional_formatting.add(f"{col}3:{col}{MAXR}", CellIsRule(operator="equal", formula=['"⚑"'], fill=RED))
don.conditional_formatting.add(f"AB3:AB{MAXR}", CellIsRule(operator="equal", formula=['"Có"'], fill=GREEN))
don.conditional_formatting.add(f"AB3:AB{MAXR}", CellIsRule(operator="equal", formula=['"Không"'], fill=RED))

# ------------------------------------------------------------------ CASE
cs = wb.create_sheet("CASE")
CASE = [("Mã case", 11, "CX"), ("Mã đơn", 11, "CX"), ("Loại", 20, "CX"), ("Kênh", 12, "CX"), ("Mở lúc", 14, "CX"),
        ("Phản hồi đầu lúc", 14, "CX"), ("Đủ bằng chứng", 11, "CX"), ("Giải pháp", 16, "CX"), ("Đền bù / hoàn (A$)", 10, "CX"),
        ("Cấp duyệt cần", 9, "SYS"), ("Người duyệt", 9, "OPS"), ("Đúng cấp?", 8, "SYS"), ("Thực thi xong", 11, "SRC"),
        ("Ngày hoàn tiền", 11, "FIN"), ("Đóng case", 11, "CX"), ("Nguyên nhân gốc", 11, "OPS"), ("Thu hồi từ", 12, "SRC"),
        ("Số thu hồi (A$)", 10, "FIN"), ("Chi phí ròng (A$)", 10, "SYS"), ("Giờ phản hồi", 9, "SYS"), ("Số ngày mở", 9, "SYS"),
        ("⚑ Phản hồi >24h", 10, "SYS"), ("⚑ Mở >7 ngày", 10, "SYS"), ("⚑ Duyệt sai cấp", 10, "SYS"),
        ("Ghi chú (không ghi dữ liệu cá nhân)", 36, "CX")]
header(cs, CASE, "SOP-B · Case hậu mãi — mỗi case một dòng · đóng case phải có nguyên nhân gốc")
CASEF = {
    "J": '=ARRAYFORMULA(IF(I3:I="","",IF((I3:I>150)+(H3:H="R4 Hoàn 100%"),"CEO",IF(I3:I>50,"OPS","CX"))))',
    "L": '=ARRAYFORMULA(IF((J3:J="")+(K3:K=""),"",IF(IFERROR(MATCH(K3:K,DS!$H$2:$H$4,0),0)>=IFERROR(MATCH(J3:J,DS!$H$2:$H$4,0),9),"✓","✗")))',
    "S": '=ARRAYFORMULA(IF(A3:A="","",IFERROR(I3:I*1,0)-IFERROR(R3:R*1,0)))',
    "T": '=ARRAYFORMULA(IF((E3:E="")+(F3:F=""),"",ROUND((F3:F-E3:E)*24,1)))',
    "U": '=ARRAYFORMULA(IF(E3:E="","",INT(IF(O3:O="",NOW(),O3:O)-E3:E)))',
    "V": '=ARRAYFORMULA(IF(E3:E="","",IF(F3:F="",IF((NOW()-E3:E)*24>24,"⚑",""),IF((F3:F-E3:E)*24>24,"⚑",""))))',
    "W": '=ARRAYFORMULA(IF((E3:E="")+(O3:O<>""),"",IF(NOW()-E3:E>7,"⚑","")))',
    "X": '=ARRAYFORMULA(IF(L3:L="✗","⚑",""))',
}
for col, f in CASEF.items():
    cs[f"{col}3"] = f
for col, rng in (("C", "=DS!$E$2:$E$9"), ("D", "=DS!$O$2:$O$8"), ("H", "=DS!$G$2:$G$9"), ("K", "=DS!$H$2:$H$4"), ("P", "=DS!$I$2:$I$9")):
    v = DataValidation(type="list", formula1=rng, allow_blank=True)
    cs.add_data_validation(v)
    v.add(f"{col}3:{col}{MAXR}")
for col in ("V", "W", "X"):
    cs.conditional_formatting.add(f"{col}3:{col}{MAXR}", CellIsRule(operator="equal", formula=['"⚑"'], fill=RED))

# ------------------------------------------------------------------ DASHBOARD
db = wb.create_sheet("DASHBOARD")
db.sheet_view.showGridLines = False
db["B2"] = "Bảng điều khiển vận hành"
db["B2"].font = font(bold=True, size=18)
db["B3"] = '="Hôm nay: "&TEXT(TODAY(),"dd/mm/yyyy")&" · từ tab DON và CASE"'
tiles = [
    ("Đơn đang chạy", '=COUNTA(DON!A3:A)-COUNTIF(DON!T3:T,"99*")-COUNTIF(DON!T3:T,"95*")', "0", ""),
    ("Giao trong khung", '=IFERROR(COUNTIF(DON!AB3:AB,"Có")/(COUNTIF(DON!AB3:AB,"Có")+COUNTIF(DON!AB3:AB,"Không")),"–")', "0%", "≥ 90%"),
    ("Đặt NCC ≤ 24h", '=IFERROR(COUNTIF(DON!AE3:AE,"<=1")/COUNT(DON!AE3:AE),"–")', "0%", "≥ 95%"),
    ("Cờ đỏ đơn đang mở", '=COUNTIF(DON!W3:AA,"⚑")', "0", "0"),
    ("Case / đơn", '=IFERROR(COUNTA(CASE!A3:A)/COUNTA(DON!A3:A),"–")', "0.0%", "≤ 3%"),
    ("Phản hồi đầu (giờ, trung vị)", '=IFERROR(MEDIAN(CASE!T3:T),"–")', "0.0", "≤ 4"),
    ("Chargeback / đơn", '=IFERROR(COUNTIF(CASE!C3:C,"C7*")/COUNTA(DON!A3:A),"–")', "0.0%", "≤ 0,5%"),
    ("Chi phí case ròng (A$)", '=SUM(CASE!S3:S)', "#,##0", ""),
]
for k, (label, f, fmt, target) in enumerate(tiles):
    col, row = 2 + (k % 4) * 3, 5 + (k // 4) * 4
    db.cell(row=row, column=col, value=label).font = font(color="6E6E73", size=10)
    v = db.cell(row=row + 1, column=col, value=f)
    v.font = font(bold=True, size=20)
    v.number_format = fmt
    if target:
        db.cell(row=row + 2, column=col, value=f"Mục tiêu {target}").font = font(color="6E6E73", size=9)
for c in range(2, 14):
    db.column_dimensions[L(c)].width = 12
T = 16
db.column_dimensions[L(T)].width = 24


def table(top, title, labels, fn):
    db.cell(row=top, column=T, value=title).font = font(bold=True)
    for i, lab in enumerate(labels):
        db.cell(row=top + 1 + i, column=T, value=lab)
        db.cell(row=top + 1 + i, column=T + 1, value=fn(lab))
    return top + 1, top + len(labels)


def chart(title, rng, anchor, color):
    ch = BarChart()
    ch.type = "bar"
    ch.title = title
    ch.legend = None
    ch.height, ch.width = 7, 13
    ch.add_data(Reference(db, min_col=T + 1, min_row=rng[0], max_row=rng[1]), titles_from_data=False)
    ch.set_categories(Reference(db, min_col=T, min_row=rng[0], max_row=rng[1]))
    ch.series[0].graphicalProperties.solidFill = color
    ch.x_axis.scaling.orientation = "maxMin"
    db.add_chart(ch, anchor)


chart("Đơn theo trạng thái", table(5, "Đơn theo trạng thái", [s for s, _ in STATUSES], lambda s: f'=COUNTIF(DON!T3:T,"{s}")'), "B14", "6B5A45")
lt = {"S1–S2 TT→Đặt NCC": "AE", "S3 Đặt→NCC gửi": "AF", "S4 Gửi→QC đạt": "AG", "S5 QC→Xuất": "AH", "S6–S7 Xuất→Giao": "AI"}
chart("Lead time TB theo giai đoạn (ngày)", table(18, "Lead time TB (ngày)", list(lt), lambda s: f'=IFERROR(AVERAGE(DON!{lt[s]}3:{lt[s]}),0)'), "H14", "0F7C80")
flags = {"Chưa đặt NCC >24h": "DON!W", "NCC quá hẹn >2d": "DON!X", "Sắp vượt ETA": "DON!Y", "Vượt ETA": "DON!Z",
         ">7d chưa cập nhật KH": "DON!AA", "Case phản hồi >24h": "CASE!V", "Case mở >7 ngày": "CASE!W", "Case duyệt sai cấp": "CASE!X"}
chart("Cờ SLA đang bật → họp 09:00", table(25, "Cờ SLA đang bật", list(flags),
      lambda s: f'=COUNTIF({flags[s]}3:{flags[s].split("!")[1]},"⚑")'), "B29", "C0392B")
chart("Case theo loại", table(35, "Case theo loại", CASES, lambda s: f'=COUNTIF(CASE!C3:C,"{s}")'), "H29", "2F6FB0")
chart("Nguyên nhân gốc", table(45, "Case theo nguyên nhân gốc", RCS, lambda s: f'=COUNTIF(CASE!P3:P,"{s}")'), "B44", "2E7D4F")

# ------------------------------------------------------------------ registry tabs (from site_map.py)
def plain(ws_title, cols, rows, note):
    ws = wb.create_sheet(ws_title)
    header(ws, [(h, w, "INK") for h, w in cols], note)
    ws.freeze_panes = "A3"
    for r, row in enumerate(rows, start=3):
        for i, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


WHO = {
    "order": ("Tab DON", "Chủ trì từng giai đoạn (CX, SRC, LOG)", "Mọi vị trí (mã đơn, không dữ liệu cá nhân)"),
    "case": ("Tab CASE", "CX mở/đóng · người duyệt ghi tên · FIN ghi ngày hoàn", "CX, OPS, FIN, CEO"),
    "sku": ("Tab MASTER_SKU", "SRC chốt thông số · MKT dùng để đăng", "Mọi vị trí (không có danh tính NCC)"),
    "supplier": ("Drive: thư mục NCC / PO / QC (hạn chế quyền)", "SRC", "SRC, OPS, CEO, FIN"),
    "shipment": ("Drive: thư mục lô hàng theo container", "LOG", "LOG, OPS, CEO, FIN"),
    "cost": ("File chi phí FIN riêng · KHÔNG đặt trong sheet này", "FIN", "FIN, CEO (OPS xem phần cần)"),
    "promise": ("Sổ tay: Customer-facing Policy + mẫu email xác nhận", "CEO duyệt · OPS sửa", "Công khai"),
    "listing": ("Shopify, đăng từ MASTER_SKU", "MKT", "Công khai sau khi duyệt"),
}
plain("SOURCE_OF_TRUTH", [("Thông tin", 30), ("Bản ghi gốc", 34), ("Ở đâu", 44), ("Ai cập nhật", 34), ("Ai được xem", 34)],
      [(i, s, *WHO[k]) for k, (i, s, _l) in SM.SOT.items()],
      "Source of truth: mỗi loại thông tin có MỘT bản ghi gốc. Sổ tay mô tả cách làm; số liệu sống nằm ở đây.")
sku = plain("MASTER_SKU", [(h, 12) for h in [
    "Mã SP", "SKU biến thể", "Nhóm chính", "Nhóm con", "Tên tiếng Anh", "Màu chuẩn (mã)", "Chất liệu từng bộ phận", "Lớp phủ",
    "W (cm)", "D (cm)", "H (cm)", "Dung sai đăng web", "Số kiện", "KT + cân nặng từng kiện", "Lắp (phút)", "Cảnh báo đổ ngã (Y/N)",
    "Gỗ / tre / mây (Y/N)", "Hướng dẫn lắp EN (link)", "Quyền dùng ảnh (bằng chứng)", "Hạng A/B/C", "Đạt 7 tiêu chí (Y/N)",
    "Mã NCC chính (nội bộ)", "Mã NCC dự phòng (nội bộ)", "Hàng mẫu đã kiểm (ngày)", "Trạng thái listing", "Người duyệt"]], [],
    "Master SKU: nguồn duy nhất cho thông số đăng web, nhãn, chứng từ. Chỉ ghi mã NCC nội bộ, không ghi tên/link NCC.")
for col, opts in (("T", '"A,B,C"'), ("U", '"Y,N"'), ("P", '"Y,N"'), ("Q", '"Y,N"'), ("Y", '"Nháp,Duyệt,Đăng,Tạm ngừng"')):
    v = DataValidation(type="list", formula1=opts, allow_blank=True)
    sku.add_data_validation(v)
    v.add(f"{col}3:{col}{MAXR}")
blk = plain("LAUNCH_BLOCKERS", [("Mã", 10), ("Việc chưa xong", 28), ("Mô tả", 60), ("Chủ", 12), ("Trạng thái", 12), ("Ngày đóng · bằng chứng", 30)],
            [(k.upper(), t, d, o, "Mở", "") for k, (t, d, o, _h) in SM.BLOCKERS.items()],
            "Launch blockers: phải đóng trước khi đăng chính sách hoặc chạy quảng cáo.")
v = DataValidation(type="list", formula1='"Mở,Đang xử lý,Đã đóng"', allow_blank=True)
blk.add_data_validation(v)
v.add("E3:E50")
blk.conditional_formatting.add("E3:E50", CellIsRule(operator="equal", formula=['"Mở"'], fill=RED))
blk.conditional_formatting.add("E3:E50", CellIsRule(operator="equal", formula=['"Đã đóng"'], fill=GREEN))

hd = wb.create_sheet("HUONG_DAN")
hd.column_dimensions["B"].width = 26
hd.column_dimensions["C"].width = 100
rows = [("[BRAND] Source of Truth", ""),
        ("Sổ tay", "https://doanngocmyy.github.io/brand-operating-handbook/docs/index.html"),
        ("Ô trắng", "Nhập tay. Chỉ nhập ngày khi đã có BẰNG CHỨNG (mã PO, ảnh QC, lượt quét thật, POD)."),
        ("Cột tiêu đề xám", "Tự tính bằng một công thức ở dòng 3 (ARRAYFORMULA). Không gõ vào các cột này."),
        ("Màu tiêu đề", "Vị trí nhập: xanh dương CX · xanh lá SRC · xanh ngọc LOG · nâu OPS · tím FIN · xám tự tính."),
        ("Họp 09:00", "DASHBOARD → “Cờ SLA đang bật”. Lọc từng cột ⚑ ở DON/CASE, giao người xử lý trong ngày."),
        ("Thứ Sáu", "OPS chụp DASHBOARD vào báo cáo tuần. Đếm case theo nguyên nhân gốc."),
        ("Dữ liệu", "KHÔNG ghi tên, email, SĐT, địa chỉ khách. Mã đơn là khóa nối. Chi phí FIN và danh tính NCC để file riêng."),
        ("Phân quyền", "Khi có nhiều người: chia sẻ theo vị trí, khóa (Protect) các cột xám và tab DS.")]
for i, (a, b) in enumerate(rows, start=2):
    hd.cell(row=i, column=2, value=a).font = font(bold=True, size=16 if i == 2 else 10)
    hd.cell(row=i, column=3, value=b).alignment = Alignment(wrap_text=True)

wb._sheets = [wb[n] for n in ("HUONG_DAN", "SOURCE_OF_TRUTH", "DASHBOARD", "DON", "CASE", "MASTER_SKU", "LAUNCH_BLOCKERS", "DS")]
wb.active = 1
wb.save(OUT)
print("saved", OUT)
